from flask import Flask, render_template, request
from werkzeug.utils import secure_filename
from docx import Document
from openpyxl import load_workbook
import csv
import io
import os
import difflib
import zlib

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024  # 10 MB

ALLOWED = {"docx", "csv", "xlsx"}

def get_extension(filename):
    return filename.rsplit(".", 1)[1].lower() if "." in filename else ""

def read_docx(file):
    doc = Document(file)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)

def read_csv(file):
    raw = file.read()
    text = raw.decode("utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(text)))
    return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)

def read_xlsx(file):
    wb = load_workbook(file, read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"[Sheet: {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            values = ["" if v is None else str(v).strip() for v in row]
            if any(values):
                lines.append(" | ".join(values))
    wb.close()
    return "\n".join(lines)

def read_file(file):
    ext = get_extension(file.filename)
    if ext not in ALLOWED:
        raise ValueError("Only .docx, .csv and .xlsx files are supported.")
    if ext == "docx":
        return read_docx(file)
    if ext == "csv":
        return read_csv(file)
    return read_xlsx(file)

def crc32(text):
    return format(zlib.crc32(text.encode("utf-8")) & 0xffffffff, "08X")

def make_diff(a, b):
    diff = list(difflib.unified_diff(
        a.splitlines(), b.splitlines(),
        fromfile="Original", tofile="Received", lineterm=""
    ))
    return "\n".join(diff[:80])

@app.route("/", methods=["GET", "POST"])
def index():
    result = None
    error = None

    if request.method == "POST":
        original_file = request.files.get("original_file")
        received_file = request.files.get("received_file")

        if not original_file or not original_file.filename or not received_file or not received_file.filename:
            error = "Please upload both the original and received files."
        else:
            try:
                original = read_file(original_file)
                received = read_file(received_file)

                sender_crc = crc32(original)
                receiver_crc = crc32(received)
                accepted = sender_crc == receiver_crc

                result = {
                    "original_name": secure_filename(original_file.filename),
                    "received_name": secure_filename(received_file.filename),
                    "sender_crc": sender_crc,
                    "receiver_crc": receiver_crc,
                    "accepted": accepted,
                    "diff": make_diff(original, received) if not accepted else ""
                }
            except Exception as exc:
                error = f"Could not read the files: {exc}"

    return render_template("index.html", result=result, error=error)

if __name__ == "__main__":
    app.run(debug=True)
